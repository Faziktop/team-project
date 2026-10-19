"""Тесты усиленной безопасности: вход, редиректы, заголовки, CSP, права, экспорт, валидация."""
import re
from datetime import date, timedelta
from pathlib import Path

import pytest

from app import create_app
from app.config import is_insecure_secret_key
from app.constants import LOGIN_MAX_FAILED_ATTEMPTS, ROLE_LOCATION_ADMIN, ROLE_HOST
from app.models import User, PromoCode, Booking
from app.security import is_safe_redirect_url, normalize_phone, sanitize_for_log, parse_iso_date
from app.services import booking_service, report_service
from tests.conftest import login

TEMPLATES = Path(__file__).resolve().parent.parent / "app" / "templates"


# ------------------------------------------------------------ open redirect
@pytest.mark.parametrize("url", [
    "//evil.com", "/\\evil.com", "\\\\evil.com", "\\evil.com", "https://evil.com", "http://evil.com/x",
    "javascript:alert(1)", "/\t/evil.com", "/\n/evil.com", "///evil.com", "", None, "evil.com", "/ok\r\nSet-Cookie: a=b",
])
def test_unsafe_redirect_targets_are_rejected(url):
    assert is_safe_redirect_url(url) is False


@pytest.mark.parametrize("url", ["/", "/admin/", "/schedule/?date=2026-10-01", "/widget/quest/3"])
def test_safe_redirect_targets_are_accepted(url):
    assert is_safe_redirect_url(url) is True


def test_login_ignores_evil_next(client, two_branches):
    a = two_branches["a"]
    r = client.post("/auth/login?next=/%5Cevil.com", data={"email": a["admin"].email, "password": a["admin_pw"]})
    assert r.status_code == 302
    assert "evil.com" not in r.headers["Location"]


def test_login_follows_safe_next(client, two_branches):
    a = two_branches["a"]
    r = client.post("/auth/login?next=/schedule/", data={"email": a["admin"].email, "password": a["admin_pw"]})
    assert r.headers["Location"].endswith("/schedule/")


# ------------------------------------------------------------ logout
def test_logout_requires_post(client, two_branches):
    a = two_branches["a"]
    login(client, a["admin"].email, a["admin_pw"])
    assert client.get("/auth/logout").status_code == 405
    assert client.get("/admin/").status_code == 200  # сессия жива
    assert client.post("/auth/logout").status_code == 302
    assert client.get("/admin/", follow_redirects=False).status_code == 302  # снова на логин


def test_logout_csrf_enforced():
    app = create_app("testing")
    app.config["WTF_CSRF_ENABLED"] = True
    with app.app_context():
        from app.extensions import db
        db.create_all()
        u = User(role=ROLE_LOCATION_ADMIN, full_name="A", email="a@example.com")
        u.set_password("Passw0rd!x")
        db.session.add(u)
        db.session.commit()
        c = app.test_client()
        with c.session_transaction() as sess:
            sess["_user_id"] = str(u.id)
            sess["_fresh"] = True
        assert c.post("/auth/logout").status_code == 400  # нет CSRF-токена
        db.drop_all()


# ------------------------------------------------------------ блокировка при переборе
def test_account_locks_after_repeated_failures(client, db, two_branches):
    a = two_branches["a"]
    for _ in range(LOGIN_MAX_FAILED_ATTEMPTS):
        login(client, a["admin"].email, "неверный-пароль1")
    user = db.session.get(User, a["admin"].id)
    assert user.is_locked()
    # даже с верным паролем во время блокировки войти нельзя
    r = login(client, a["admin"].email, a["admin_pw"])
    assert r.status_code == 200  # форма перерисована, без редиректа
    assert client.get("/admin/", follow_redirects=False).status_code == 302


def test_lock_expires_and_success_resets_counter(client, db, two_branches):
    from datetime import datetime, timedelta as td, timezone
    a = two_branches["a"]
    for _ in range(LOGIN_MAX_FAILED_ATTEMPTS):
        login(client, a["admin"].email, "неверный-пароль1")
    user = db.session.get(User, a["admin"].id)
    user.locked_until = datetime.now(timezone.utc).replace(tzinfo=None) - td(minutes=1)
    db.session.commit()
    assert login(client, a["admin"].email, a["admin_pw"]).status_code == 302
    user = db.session.get(User, a["admin"].id)
    assert user.failed_login_count == 0 and user.locked_until is None


def test_failed_logins_do_not_reveal_whether_email_exists(client, two_branches):
    a = two_branches["a"]
    r1 = login(client, "nobody@example.com", "какой-то-пароль1")
    r2 = login(client, a["admin"].email, "какой-то-пароль1")
    assert r1.status_code == r2.status_code == 200
    assert r1.get_data() == r2.get_data() or "Неверный email или пароль" in r1.get_data(as_text=True)


def test_unknown_email_failure_is_audited_without_crash(client, db):
    assert login(client, "ghost@example.com", "xxxxxxxx1").status_code == 200


def test_session_is_regenerated_on_login(client, two_branches):
    a = two_branches["a"]
    with client.session_transaction() as sess:
        sess["attacker_value"] = "planted"
    login(client, a["admin"].email, a["admin_pw"])
    with client.session_transaction() as sess:
        assert "attacker_value" not in sess


# ------------------------------------------------------------ пароли
@pytest.mark.parametrize("pw", ["short1", "onlyletters", "12345678", "password123", "test1234"])
def test_weak_passwords_rejected_on_register(client, pw):
    r = client.post("/auth/register", data={"full_name": "Тест", "phone": "+79001112233", "password": pw})
    assert r.status_code == 200
    assert User.query.filter_by(phone="+79001112233").first() is None


def test_strong_password_accepted_on_register(client):
    r = client.post("/auth/register", data={"full_name": "Тест", "phone": "+79001112233", "password": "Хороший-пароль-2026"})
    assert r.status_code == 302
    assert User.query.filter_by(phone="+79001112233").first() is not None


def test_password_equal_to_email_rejected(client):
    r = client.post("/auth/register", data={"full_name": "Т", "email": "abc12345@example.com", "phone": "+79001112244",
                                            "password": "abc12345@example.com"})
    assert User.query.filter_by(phone="+79001112244").first() is None


def test_very_long_password_rejected_not_hashed(client, two_branches):
    r = login(client, two_branches["a"]["admin"].email, "a1" * 5000)
    assert r.status_code == 200


# ------------------------------------------------------------ заголовки / CSP
def test_security_headers_present(client):
    r = client.get("/")
    csp = r.headers["Content-Security-Policy"]
    assert "default-src 'self'" in csp and "script-src 'self'" in csp
    assert "frame-ancestors 'none'" in csp and "object-src 'none'" in csp
    assert "unsafe-inline" not in csp.split("script-src")[1].split(";")[0]
    assert r.headers["X-Frame-Options"] == "DENY"
    assert r.headers["X-Content-Type-Options"] == "nosniff"
    assert "Permissions-Policy" in r.headers and "Cross-Origin-Opener-Policy" in r.headers


def test_dynamic_pages_not_cacheable_but_static_is(client):
    assert client.get("/").headers["Cache-Control"] == "no-store"
    static = client.get("/static/js/app.js")
    assert static.status_code == 200
    assert static.headers.get("Cache-Control") != "no-store"


def test_no_external_hosts_in_base_template(client):
    html = client.get("/").get_data(as_text=True)
    assert not re.search(r'(?:src|href)="https?://', html), "внешние ресурсы запрещены CSP"


def test_templates_have_no_inline_scripts_or_handlers():
    """CSP запрещает inline-JS: любой такой код в шаблоне просто перестал бы работать."""
    offenders = []
    for path in TEMPLATES.rglob("*.html"):
        text = path.read_text(encoding="utf-8")
        if re.search(r"<script(?![^>]*\bsrc=)[^>]*>", text):
            offenders.append(f"{path.name}: inline <script>")
        if re.search(r"\son(?:click|submit|change|load|error|mouseover|focus|input)\s*=", text, re.I):
            offenders.append(f"{path.name}: inline-обработчик on*=")
        if re.search(r"javascript:", text, re.I):
            offenders.append(f"{path.name}: javascript: URL")
    assert not offenders, offenders


def test_vendored_assets_exist(client):
    for url in ("/static/vendor/bootstrap/bootstrap.min.css", "/static/vendor/bootstrap/bootstrap.bundle.min.js",
                "/static/vendor/bootstrap-icons/bootstrap-icons.min.css",
                "/static/vendor/bootstrap-icons/fonts/bootstrap-icons.woff2", "/static/vendor/fonts/fonts.css"):
        assert client.get(url).status_code == 200, url


# ------------------------------------------------------------ конфигурация
@pytest.mark.parametrize("key", [None, "", "short", "dev-secret-key-change-me", "change-me-in-production",
                                 "Change-Me-" + "x" * 40, "secret" + "y" * 40, "test" + "z" * 40])
def test_insecure_secret_keys_detected(key):
    assert is_insecure_secret_key(key)


def test_strong_secret_key_is_accepted():
    import secrets
    assert not is_insecure_secret_key(secrets.token_hex(32))


def test_production_refuses_weak_secret(monkeypatch):
    monkeypatch.setattr("app.config.Config.SECRET_KEY", "change-me-in-production", raising=False)
    monkeypatch.setattr("app.config.ProductionConfig.SECRET_KEY", "change-me-in-production", raising=False)
    with pytest.raises(RuntimeError):
        create_app("production")


# ------------------------------------------------------------ обработка ввода: 400 вместо 500
@pytest.mark.parametrize("qs", ["date_from=garbage", "date_to=2026-13-45", "date_from=2026-02-30",
                                "date_from=2026-10-10&date_to=2026-10-01",
                                "date_from=2000-01-01&date_to=2026-01-01"])
def test_reports_bad_period_returns_400(client, two_branches, qs):
    a = two_branches["a"]
    login(client, a["admin"].email, a["admin_pw"])
    assert client.get(f"/reports/?{qs}").status_code == 400
    assert client.get(f"/reports/export/occupancy.csv?{qs}").status_code == 400


def test_calendar_bad_date_returns_400(client, two_branches):
    a = two_branches["a"]
    login(client, a["admin"].email, a["admin_pw"])
    assert client.get("/schedule/?date=bad").status_code == 400


def test_public_slots_api_validation(client, two_branches):
    qid = two_branches["a"]["quest"].id
    d = (date.today() + timedelta(days=1)).isoformat()
    assert client.get(f"/schedule/api/available-slots?quest_id={qid}&date_from={d}&date_to={d}").status_code == 200
    r = client.get(f"/schedule/api/available-slots?quest_id={qid}&date_from=bad&date_to={d}")
    assert r.status_code == 400 and r.is_json
    far = (date.today() + timedelta(days=400)).isoformat()
    assert client.get(f"/schedule/api/available-slots?quest_id={qid}&date_from={d}&date_to={far}").status_code == 400
    assert client.get(f"/schedule/api/available-slots?quest_id={qid}&date_from={far}&date_to={d}").status_code == 400
    assert client.get("/schedule/api/available-slots").status_code == 400


def test_export_unknown_format_and_report(client, two_branches):
    a = two_branches["a"]
    login(client, a["admin"].email, a["admin_pw"])
    assert client.get("/reports/export/occupancy.pdf").status_code == 400
    assert client.get("/reports/export/nothing.csv").status_code == 404


def test_promocode_form_rejects_bad_dates_without_500(client, db, two_branches):
    a = two_branches["a"]
    login(client, a["admin"].email, a["admin_pw"])
    base = {"code": "NEWCODE", "discount_percent": "10", "discount_amount": "0", "usage_limit": "0"}
    assert client.post("/payments/promocodes/new", data={**base, "valid_from": "garbage"}).status_code == 200
    assert client.post("/payments/promocodes/new", data={**base, "valid_from": "2026-12-01", "valid_to": "2026-01-01"}).status_code == 200
    assert PromoCode.query.filter_by(code="NEWCODE").first() is None
    r = client.post("/payments/promocodes/new", data={**base, "valid_from": "2026-01-01", "valid_to": "2026-12-01"})
    assert r.status_code == 302 and PromoCode.query.filter_by(code="NEWCODE").first()


def test_promocode_requires_some_discount(client, two_branches):
    a = two_branches["a"]
    login(client, a["admin"].email, a["admin_pw"])
    r = client.post("/payments/promocodes/new", data={"code": "EMPTYPROMO", "discount_percent": "0", "discount_amount": "0"})
    assert r.status_code == 200 and PromoCode.query.filter_by(code="EMPTYPROMO").first() is None


def test_certificate_in_past_rejected(client, two_branches):
    a = two_branches["a"]
    login(client, a["admin"].email, a["admin_pw"])
    r = client.post("/payments/certificates/new", data={"code": "OLDCERT", "initial_balance": "500", "expires_at": "2020-01-01"})
    assert r.status_code == 200


def test_tariff_time_format_validated(client, two_branches):
    a = two_branches["a"]
    login(client, a["admin"].email, a["admin_pw"])
    from app.models import Tariff
    n = Tariff.query.count()
    data = {"quest_id": a["quest"].id, "name": "Т", "price_per_slot": "500", "time_from": "9:00", "time_to": "23:59"}
    assert client.post("/admin/tariffs/new", data=data).status_code == 200
    data.update(time_from="18:00", time_to="10:00")
    assert client.post("/admin/tariffs/new", data=data).status_code == 200
    assert Tariff.query.count() == n


def test_slot_generation_form_rejects_foreign_quest(client, db, two_branches):
    a, b = two_branches["a"], two_branches["b"]
    login(client, a["admin"].email, a["admin_pw"])
    from app.models import Slot
    before = Slot.query.count()
    d = (date.today() + timedelta(days=9)).isoformat()
    r = client.post("/schedule/generate", data={"room_id": a["room"].id, "quest_id": b["quest"].id,
                                                "date_from": d, "date_to": d, "day_start": "10:00", "day_end": "12:00"})
    assert r.status_code == 200 and Slot.query.count() == before  # чужой квест не входит в choices формы
    r = client.post("/schedule/generate", data={"room_id": a["room"].id, "quest_id": a["quest"].id,
                                                "date_from": d, "date_to": d, "day_start": "10:00", "day_end": "12:00"})
    assert r.status_code == 302 and Slot.query.count() == before + 2


# ------------------------------------------------------------ экспорт: formula injection
def test_csv_export_neutralizes_formulas():
    rows = [{"name": "=HYPERLINK(\"http://evil\")", "n": 5}, {"name": "+1+1", "n": 6}, {"name": "@SUM(A1)", "n": 7},
            {"name": "-2+3", "n": 8}, {"name": "Обычное", "n": 9}]
    text = report_service.export_rows_to_csv(rows).read().decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'+1+1" in text and "'@SUM" in text and "'-2+3" in text
    assert "Обычное" in text and "'Обычное" not in text


def test_xlsx_export_neutralizes_formulas():
    from openpyxl import load_workbook
    buf = report_service.export_rows_to_xlsx([{"name": "=1+1", "n": 1}])
    ws = load_workbook(buf).active
    assert ws["A2"].value == "'=1+1" and ws["A2"].data_type != "f"


# ------------------------------------------------------------ вспомогательные функции
def test_normalize_phone_variants_match():
    assert normalize_phone("8 (900) 123-45-67") == normalize_phone("+7 900 123 45 67") == "79001234567"
    assert normalize_phone(None) == "" and normalize_phone("abc") == ""


def test_sanitize_for_log_strips_newlines():
    assert "\n" not in sanitize_for_log("a\nWARNING fake\r\nline") and len(sanitize_for_log("x" * 9999)) == 500


def test_parse_iso_date():
    assert parse_iso_date("2026-02-28") == date(2026, 2, 28) and parse_iso_date("") is None
    with pytest.raises(ValueError):
        parse_iso_date("2026-02-30")


# ------------------------------------------------------------ гостевая бронь через виджет
def _widget_post(client, slot_id, **over):
    data = {"slot_id": slot_id, "full_name": "Гость", "phone": "+79005550001", "players_count": "2"}
    data.update(over)
    return client.post(f"/booking/step2/{slot_id}", data=data)


def test_guest_booking_happy_path(client, db, two_branches):
    slot = two_branches["a"]["slots"][0]
    r = _widget_post(client, slot.id, email="guest@example.com")
    assert r.status_code == 302 and "/booking/confirmation/" in r.headers["Location"]
    guest = User.query.filter_by(phone="+79005550001").first()
    assert guest.role == "client" and guest.email == "guest@example.com"


def test_guest_cannot_attach_booking_to_staff_by_phone(client, db, two_branches):
    a = two_branches["a"]
    a["manager"].phone = "+79005559999"
    db.session.commit()
    r = _widget_post(client, a["slots"][0].id, phone="+79005559999")
    assert r.status_code == 200  # ошибка, а не редирект/500
    assert Booking.query.count() == 0


def test_guest_with_taken_email_does_not_crash(client, db, two_branches):
    a = two_branches["a"]
    r = _widget_post(client, a["slots"][0].id, email=a["admin"].email)
    assert r.status_code == 302
    guest = User.query.filter_by(phone="+79005550001").first()
    assert guest.email is None  # чужой email не присваивается


def test_widget_rejects_invalid_phone_and_email(client, two_branches):
    slot = two_branches["a"]["slots"][0]
    assert _widget_post(client, slot.id, phone="<script>").status_code == 200
    assert _widget_post(client, slot.id, email="not-an-email").status_code == 200
    assert Booking.query.count() == 0


def test_lookup_matches_normalized_phone_and_hides_others(client, db, two_branches):
    a = two_branches["a"]
    b = booking_service.create_booking(a["client"], a["slots"][0].id, 2)
    ok = client.post("/lookup", data={"code": b.confirmation_code, "phone": a["client"].phone.replace("+7", "8")})
    assert "Квест a" in ok.get_data(as_text=True)  # найдено: страница показывает квест брони
    bad = client.post("/lookup", data={"code": b.confirmation_code, "phone": "+79990000000"})
    assert "не найдено" in bad.get_data(as_text=True) and "Квест a" not in bad.get_data(as_text=True)
    empty = client.post("/lookup", data={"code": b.confirmation_code, "phone": ""})
    assert "не найдено" in empty.get_data(as_text=True)


def test_client_cannot_cancel_someone_elses_booking(client, db, two_branches):
    a, b = two_branches["a"], two_branches["b"]
    booking = booking_service.create_booking(b["client"], b["slots"][0].id, 2)
    a["client"].set_password("Clientpass1")
    a["client"].email = "ca@example.com"
    db.session.commit()
    login(client, "ca@example.com", "Clientpass1")
    assert client.post(f"/my-bookings/{booking.id}/cancel").status_code == 404


def test_inactive_quest_hidden_from_widget(client, db, two_branches):
    a = two_branches["a"]
    a["quest"].is_active = False
    db.session.commit()
    assert client.get(f"/quest/{a['quest'].id}").status_code == 404
    assert _widget_post(client, a["slots"][0].id).status_code == 302  # редирект на страницу квеста, брони нет
    assert Booking.query.count() == 0


# ------------------------------------------------------------ rate limiting
def test_login_is_rate_limited(monkeypatch):
    # limiter — глобальный объект и читает настройки при каждом init_app(): включаем его
    # через конфиг ДО создания приложения; следующее create_app() снова применит обычный конфиг.
    monkeypatch.setattr("app.config.TestingConfig.RATELIMIT_ENABLED", True)
    app = create_app("testing")
    with app.app_context():
        from app.extensions import db
        db.create_all()
        c = app.test_client()
        codes = [c.post("/auth/login", data={"email": "x@example.com", "password": "Passw0rd1"}).status_code
                 for _ in range(14)]
        db.drop_all()
    assert 429 in codes and codes[0] != 429
